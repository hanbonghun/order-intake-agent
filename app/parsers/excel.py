from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def excel_to_text(path: Path) -> str:
    """시트의 비어 있지 않은 행을 ' | '로 이어 붙인다. 표 구조는 LLM이 해석한다."""
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            header = f"[시트: {ws.title}]\n" if len(wb.worksheets) > 1 else ""
            blocks.append(header + "\n".join(rows))
    wb.close()
    return "\n\n".join(blocks)
