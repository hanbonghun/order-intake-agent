"""평가용 발주 샘플 생성 → data/samples/{case_id}/

정답(expected.json)을 먼저 만들고, 그 정답을 이메일·PDF·엑셀로 렌더링한다.
깨끗한 발주는 템플릿 렌더링(재현 가능, 비용 0), 지저분한 발주는 2단계에서 LLM 변형 + 직접 작성.

    uv run python -m scripts.generate_samples
"""
import json
import random
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

from scripts.seed_data import customers_by_id, load, products_by_sku, profiles, to_base_quantity

ROOT = Path(__file__).resolve().parent.parent
SAMPLES_DIR = ROOT / "data" / "samples"
FONT_DIR = ROOT / "data" / "fonts"

GENERATED_CATEGORIES = {"clean_email": 15, "pdf": 10, "excel": 5}
FIRST_RECEIVED = date(2026, 10, 5)
WEEKDAYS = "월화수목금토일"
# 기본 배송지가 없는 거래처가 본문에 적는 주소
FALLBACK_ADDRESS = "서울특별시 중구 을지로 100, 1층"


# ---------------------------------------------------------------------------
# 정답 생성
# ---------------------------------------------------------------------------

class Ctx:
    def __init__(self, seed: int):
        self.rng = random.Random(seed)
        self.products = products_by_sku()
        self.customers = customers_by_id()
        self.profiles = profiles()
        self.public_aliases: dict[str, list[str]] = {}
        self.customer_aliases: dict[tuple[int, str], list[str]] = {}
        for a in load("product_aliases"):
            if a["customer_id"] is None:
                self.public_aliases.setdefault(a["sku"], []).append(a["alias"])
            else:
                self.customer_aliases.setdefault((a["customer_id"], a["sku"]), []).append(a["alias"])


def display_name(ctx: Ctx, cid: int, sku: str) -> str:
    """정식 명칭 또는 등록된 별칭 중 하나. 1단계 정확 일치로 풀 수 있는 표현만 쓴다."""
    names = [ctx.products[sku]["name"]] * 3
    names += ctx.public_aliases.get(sku, []) + ctx.customer_aliases.get((cid, sku), [])
    return ctx.rng.choice(names)


def build_order(ctx: Ctx, cid: int, received: datetime) -> dict:
    rng = ctx.rng
    delivery = received.date() + timedelta(days=rng.randint(1, 3))
    if delivery.weekday() == 6:  # 일요일 배송 없음
        delivery += timedelta(days=1)

    customer = ctx.customers[cid]
    profile = ctx.profiles[cid]
    items = []
    for sku, unit, lo, hi in rng.sample(profile, k=rng.randint(2, min(5, len(profile)))):
        product = ctx.products[sku]
        qty = rng.randint(lo, hi)
        base_qty = to_base_quantity(product, qty, unit)
        if base_qty > product["stock"]:  # 깨끗한 발주는 재고 부족이 없어야 한다
            continue
        items.append({
            "sku": sku,
            "quantity": qty,
            "unit": unit,
            "base_quantity": base_qty,
            "display": display_name(ctx, cid, sku),
        })

    return {
        "is_order": True,
        "decision": "create",
        "customer_id": cid,
        "delivery_date": delivery.isoformat(),
        "delivery_address": customer["default_address"] or FALLBACK_ADDRESS,
        "address_in_text": customer["default_address"] is None,
        "items": items,
    }


# ---------------------------------------------------------------------------
# 표현 헬퍼
# ---------------------------------------------------------------------------

def fmt_date(rng: random.Random, iso: str) -> str:
    d = date.fromisoformat(iso)
    return rng.choice([
        f"{d.month}월 {d.day}일({WEEKDAYS[d.weekday()]})",
        f"{d.month}/{d.day}",
        d.isoformat(),
        f"{d.month}월 {d.day}일",
        f"{d.month}/{d.day}({WEEKDAYS[d.weekday()]})",
    ])


def fmt_unit(rng: random.Random, unit: str) -> str:
    return rng.choice(["박스", "박스", "box"]) if unit == "박스" else unit


def fmt_qty(rng: random.Random, item: dict) -> str:
    sep = rng.choice(["", " "])
    return f"{item['quantity']}{sep}{fmt_unit(rng, item['unit'])}"


# ---------------------------------------------------------------------------
# 이메일 본문 렌더링
# ---------------------------------------------------------------------------

def render_email(ctx: Ctx, order: dict) -> tuple[str, str]:
    rng = ctx.rng
    name = ctx.customers[order["customer_id"]]["name"]
    d = fmt_date(rng, order["delivery_date"])
    addr = f"배송지: {order['delivery_address']}\n" if order["address_in_text"] else ""
    lines = [(it["display"], fmt_qty(rng, it)) for it in order["items"]]
    style = rng.randrange(4)

    if style == 0:
        subject = f"[{name}] 발주 요청"
        items = "\n".join(f"- {n} {q}" for n, q in lines)
        body = f"안녕하세요, {name}입니다.\n{d} 납품 부탁드립니다.\n{addr}\n{items}\n\n감사합니다.\n{name} 드림"
    elif style == 1:
        subject = "발주드립니다"
        items = ", ".join(f"{n} {q}" for n, q in lines)
        body = f"{d}까지 {items} 보내주세요.\n{addr}-{name}"
    elif style == 2:
        subject = f"{name} 주문 ({d})"
        items = "\n".join(f"{i}. {n} / {q}" for i, (n, q) in enumerate(lines, 1))
        body = f"신선마켓 담당자님께,\n\n아래 품목 발주합니다.\n\n{items}\n\n납품일: {d}\n{addr}\n{name}"
    else:
        subject = f"[발주] {name}"
        items = "\n".join(f"{n}\t{q}" for n, q in lines)
        body = f"품목\t수량\n{items}\n\n납품 희망일 {d}\n{addr}잘 부탁드립니다."
    return subject, body


# ---------------------------------------------------------------------------
# PDF 렌더링 (양식 3종)
# ---------------------------------------------------------------------------

def _pdf() -> FPDF:
    pdf = FPDF()
    pdf.add_font("Nanum", fname=str(FONT_DIR / "NanumGothic-Regular.ttf"))
    pdf.add_font("Nanum", style="B", fname=str(FONT_DIR / "NanumGothic-Bold.ttf"))
    pdf.add_page()
    pdf.set_font("Nanum", size=11)
    return pdf


def _pdf_address(ctx: Ctx, order: dict) -> str:
    if order["address_in_text"]:
        return order["delivery_address"]
    return ctx.rng.choice([order["delivery_address"], "기본 배송지"])


def render_pdf(ctx: Ctx, order: dict, received: datetime, path: Path) -> None:
    rng = ctx.rng
    name = ctx.customers[order["customer_id"]]["name"]
    style = rng.randrange(3)
    pdf = _pdf()

    if style == 0:  # 공급업체(신선마켓) 표준 양식
        pdf.set_font("Nanum", "B", 18)
        pdf.cell(0, 12, "신선마켓 발주서", align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Nanum", size=11)
        for label, value in [
            ("거래처명", name),
            ("발주일", received.date().isoformat()),
            ("납품요청일", order["delivery_date"]),
            ("배송지", _pdf_address(ctx, order)),
        ]:
            pdf.cell(0, 8, f"{label}: {value}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)
        with pdf.table(col_widths=(10, 60, 15, 15, 30)) as table:
            table.row(["No", "품목명", "단위", "수량", "비고"])
            for i, it in enumerate(order["items"], 1):
                table.row([str(i), it["display"], fmt_unit(rng, it["unit"]), str(it["quantity"]), ""])

    elif style == 1:  # 거래처 자체 양식 A
        pdf.set_font("Nanum", "B", 20)
        pdf.cell(0, 14, "발   주   서", align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Nanum", size=11)
        pdf.cell(0, 8, f"수신: 신선마켓 귀중", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 8, f"발신: {name}", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 8, "아래와 같이 발주합니다.", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        with pdf.table(col_widths=(70, 30)) as table:
            table.row(["품명", "수량"])
            for it in order["items"]:
                table.row([it["display"], fmt_qty(rng, it)])
        pdf.ln(4)
        pdf.cell(0, 8, f"납기: {fmt_date(rng, order['delivery_date'])}", new_x="LMARGIN", new_y="NEXT")
        if order["address_in_text"]:
            pdf.cell(0, 8, f"납품 장소: {order['delivery_address']}", new_x="LMARGIN", new_y="NEXT")

    else:  # 거래처 자체 양식 B (영문 혼용)
        pdf.set_font("Nanum", "B", 16)
        pdf.cell(0, 12, "ORDER SHEET", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Nanum", size=11)
        pdf.cell(0, 8, f"Store: {name}", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 8, f"Delivery: {order['delivery_date']}", new_x="LMARGIN", new_y="NEXT")
        if order["address_in_text"]:
            pdf.cell(0, 8, f"Address: {order['delivery_address']}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        with pdf.table(col_widths=(60, 20, 20)) as table:
            table.row(["Item", "Qty", "Unit"])
            for it in order["items"]:
                table.row([it["display"], str(it["quantity"]), fmt_unit(rng, it["unit"])])

    pdf.output(str(path))


# ---------------------------------------------------------------------------
# 엑셀 렌더링 (양식 2종)
# ---------------------------------------------------------------------------

def render_excel(ctx: Ctx, order: dict, path: Path) -> None:
    rng = ctx.rng
    name = ctx.customers[order["customer_id"]]["name"]
    delivery = date.fromisoformat(order["delivery_date"])
    wb = Workbook()
    ws = wb.active
    ws.title = "발주"

    if rng.randrange(2) == 0:  # 상단 정보 + 품목 표
        ws["A1"] = "발주서"
        ws["A3"], ws["B3"] = "거래처", name
        ws["A4"], ws["B4"] = "납품일", delivery
        ws["B4"].number_format = "yyyy-mm-dd"
        if order["address_in_text"]:
            ws["A5"], ws["B5"] = "배송지", order["delivery_address"]
        ws.append([])
        ws.append(["품목", "수량", "단위"])
        for it in order["items"]:
            ws.append([it["display"], it["quantity"], fmt_unit(rng, it["unit"])])
    else:  # 행마다 납품일자가 반복되는 목록형
        ws.append(["납품일자", "품목명", "수량", "단위", "비고"])
        for it in order["items"]:
            ws.append([f"{delivery.month}/{delivery.day}", it["display"], it["quantity"], fmt_unit(rng, it["unit"]), ""])
        ws.append([])
        ws.append(["주문처", name])
        if order["address_in_text"]:
            ws.append(["배송지", order["delivery_address"]])

    wb.save(path)


ATTACHMENT_BODIES = [
    "발주서 첨부합니다. 확인 부탁드립니다.",
    "안녕하세요. 이번 주 발주서 보내드립니다.",
    "첨부파일 확인 부탁드려요~",
    "발주서 송부드립니다.\n감사합니다.",
]


# ---------------------------------------------------------------------------
# 저장
# ---------------------------------------------------------------------------

def write_case(case_id: str, category: str, description: str, ctx: Ctx, order: dict,
               received: datetime, subject: str, body: str, attachments: list[str], case_dir: Path) -> None:
    customer = ctx.customers[order["customer_id"]]
    domain = customer["email"].split("@")[1]
    input_doc = {
        "message_id": f"<{case_id}@{domain}>",
        "in_reply_to": None,
        "sender_email": customer["email"],
        "subject": subject,
        "body": body,
        "attachments": attachments,
        "received_at": received.isoformat(),
    }
    expected = {
        "case_id": case_id,
        "category": category,
        "description": description,
        "expected": {
            "is_order": order["is_order"],
            "decision": order["decision"],
            "customer_id": order["customer_id"],
            "delivery_date": order["delivery_date"],
            "delivery_address": order["delivery_address"],
            "items": [
                {k: it[k] for k in ("sku", "quantity", "unit", "base_quantity")} | {"raw_text": it["display"]}
                for it in order["items"]
            ],
        },
    }
    (case_dir / "input.json").write_text(json.dumps(input_doc, ensure_ascii=False, indent=2), encoding="utf-8")
    (case_dir / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")


def main(seed: int = 42):
    ctx = Ctx(seed)
    customer_ids = sorted(ctx.customers)
    n = 0
    for category, count in GENERATED_CATEGORIES.items():
        for i in range(1, count + 1):
            case_id = f"{category}_{i:03d}"
            case_dir = SAMPLES_DIR / case_id
            if case_dir.exists():
                shutil.rmtree(case_dir)
            case_dir.mkdir(parents=True)

            cid = customer_ids[n % len(customer_ids)]
            received = datetime.combine(
                FIRST_RECEIVED + timedelta(days=n % 10),
                datetime.min.time(),
            ).replace(hour=ctx.rng.randint(7, 17), minute=ctx.rng.randint(0, 59))
            order = build_order(ctx, cid, received)

            if category == "clean_email":
                subject, body = render_email(ctx, order)
                attachments = []
            elif category == "pdf":
                render_pdf(ctx, order, received, case_dir / "order.pdf")
                subject, body, attachments = "발주서 송부", ctx.rng.choice(ATTACHMENT_BODIES), ["order.pdf"]
            else:
                render_excel(ctx, order, case_dir / "order.xlsx")
                subject, body, attachments = "발주서(엑셀)", ctx.rng.choice(ATTACHMENT_BODIES), ["order.xlsx"]

            write_case(case_id, category, "깨끗한 발주: 등록된 명칭·별칭, 단위·납기 명시",
                       ctx, order, received, subject, body, attachments, case_dir)
            n += 1
    print(f"{n}건 생성 → {SAMPLES_DIR}")


if __name__ == "__main__":
    main()
